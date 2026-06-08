#!/usr/bin/env python3
# encoding: utf-8
'''
asap.prepareJSONInput -- Create a JSON input file for ASAP from a multifasta, GenBank, or Excel spreadsheet
'''

import sys
import os
import re
import argparse
import logging
import skbio.io
from skbio import DNA
from openpyxl import load_workbook

# We keep assayInfo because it defines the data structures for the JSON
from asap import assayInfo
from asap import __version__

__updated__ = '2026-04-02' 
__date__ = '2015-08-03'

DEBUG = 1
TESTRUN = 0
PROFILE = 0

PRESENCE_ABSENCE = 10
GENE_VARIANT = 20

def _process_genbank(gb_file):
    """
    Simplified GenBank parser. 
    Uses the FILENAME as the assay name and matches FASTA logic for ASAP validation.
    """
    return_list = []
    
    # Extract the filename without the extension to use as the name
    base_name = os.path.splitext(os.path.basename(gb_file))[0]
    
    for seq in skbio.io.registry.read(gb_file, format='genbank', constructor=DNA):
        # 1. Use the filename as the assay name
        assay_name = base_name
        
        # 2. Extract and clean sequence
        full_seq_str = _clean_seq(str(seq))
        
        # 3. Create Amplicon
        amplicon = assayInfo.Amplicon(sequence=full_seq_str)
        
        # 4. Wrap it in a Target and Assay
        # Using 'species ID' and 'presence/absence' to satisfy ASAP strict validation
        target = assayInfo.Target(function='species ID', amplicon=amplicon)
        assay = assayInfo.Assay(name=assay_name, assay_type='presence/absence', target=target)
        
        return_list.append(assay)
        
    return return_list

def _process_fasta(fasta, fasta_type, message=None):
    return_list = []
    for seq in skbio.io.registry.read(fasta, format='fasta', constructor=DNA):
        significance = assayInfo.Significance(seq.metadata['description']) if seq.metadata['description'] else message
        amplicon = assayInfo.Amplicon(sequence=_clean_seq(str(seq)), significance=significance)
        if fasta_type == GENE_VARIANT:
            amplicon.variant_name = seq.metadata['id']
            return_list.append(amplicon)
        else:
            target = assayInfo.Target(function='species ID', amplicon=amplicon)
            assay = assayInfo.Assay(name=seq.metadata['id'], assay_type='presence/absence', target=target)
            return_list.append(assay)
    return return_list

def _process_fasta_single(fasta):
    for seq in skbio.io.registry.read(fasta, format='fasta', constructor=DNA):
        amplicon = assayInfo.Amplicon(sequence=_clean_seq(str(seq)))
    return amplicon

def _clean_seq(sequence):
    return_seq = sequence.upper()
    return_seq = re.sub('[-|_]', '', return_seq)
    return return_seq

def _clean_str(string):
    return re.sub(' ', '_', string) if string else None

def _strip(string):
    if string or string == 0:
        return str(string).strip() if isinstance(string, str) else str(string)
    return None

def _isNT(sequence, positions):
    size = 0
    for token in positions.split(','):
        m = re.search(r"(\d*)-(\d*)", token)
        if m:
            size += int(m.group(2)) - int(m.group(1)) + 1
        else:
            size += 1
    return size == len(sequence)

def main(argv=None):
    if argv is None:
        # sys.argv[0] is the script name; sys.argv[1:] are the actual flags
        argv = sys.argv[1:]
    
    program_version = "v%s" % __version__
    program_build_date = str(__updated__)
    program_version_message = 'prepareJSONInput %s (%s)' % (program_version, program_build_date)
    program_shortdesc = "ASAP JSON Preparer - Standalone Version"

    try:
        # Setup argument parser
        parser = argparse.ArgumentParser(description=program_shortdesc, formatter_class=argparse.RawDescriptionHelpFormatter)
        required_group = parser.add_argument_group("required arguments")
        exclusive_group = required_group.add_mutually_exclusive_group(required=True)
        
        exclusive_group.add_argument("-f", "--fasta", metavar="FILE", help="fasta file containing amplicon sequences.")
        exclusive_group.add_argument("-x", "--excel", metavar="FILE", help="Excel file of assay data.")
        # Added nargs='+' to support multiple files at once
        exclusive_group.add_argument("-g", "--gbb", metavar="FILE", nargs='+', help="One or more GenBank files containing genome data.")
        
        required_group.add_argument("-o", "--out", metavar="FILE", required=True, help="output JSON file to write. [REQUIRED]")
        parser.add_argument("-w", "--worksheet", help="Excel worksheet to use.")
        parser.add_argument('-v', '--version', action='version', version=program_version_message)

        # Parse arguments from the passed argv
        args = parser.parse_args(argv)

        fasta_file = args.fasta
        excel_file = args.excel
        gb_files = args.gbb
        out_file = args.out
        worksheet = args.worksheet

        assay_list = []

        # 3. Processing logic
        if fasta_file:
            assay_list = _process_fasta(fasta_file, PRESENCE_ABSENCE)
            
        elif gb_files:
            # Iterate through the list of GenBank files provided
            for gb_file in gb_files:
                assay_list.extend(_process_genbank(gb_file))
                
        elif excel_file:
            wb = load_workbook(excel_file, read_only=True)
            ws = wb.active if not worksheet else wb[worksheet]
            
            amplicon = None
            target = None
            assay = None
            
            for row in ws.iter_rows(min_row=3):
                if _strip(row[0].value):  # Start a new Assay
                    if assay:
                        assay_list.append(assay)
                        target = None
                        amplicon = None
                    assay = assayInfo.Assay(name=_clean_str(_strip(row[0].value)), assay_type=_strip(row[1].value))

                if _strip(row[18].value):
                    significance = assayInfo.Significance(message=_strip(row[17].value), resistance=_strip(row[18].value))
                else:
                    significance = assayInfo.Significance(message=_strip(row[17].value))

                element = None
                if _strip(row[14].value):  # Significance attaches to a Region of Interest
                    sequence = _strip(row[15].value)
                    positions = _strip(row[14].value)
                    if _isNT(sequence, positions):
                        element = assayInfo.RegionOfInterest(position_range=positions, nt_sequence=sequence, mutations=_strip(row[16].value), name=_strip(row[13].value), significance=significance)
                    else:
                        element = assayInfo.RegionOfInterest(position_range=positions, aa_sequence=sequence, mutations=_strip(row[16].value), name=_strip(row[13].value), significance=significance)
                elif _strip(row[10].value):  # Significance attaches to a SNP
                    element = assayInfo.SNP(position=_strip(row[10].value), reference=_strip(row[11].value), variant=_strip(row[12].value), name=_strip(row[9].value), significance=significance)

                if _strip(row[8].value):  # New Amplicon sequence on this row
                    if os.path.isfile(_strip(row[8].value)):
                        if assay.assay_type == "gene variant":
                            amplicon = _process_fasta(_strip(row[8].value), GENE_VARIANT, significance)
                        else:
                            amplicon = _process_fasta_single(_strip(row[8].value))
                            if element:
                                amplicon.add_SNP(element) if isinstance(element, assayInfo.SNP) else amplicon.add_ROI(element)
                            else:
                                amplicon.significance = significance
                    else:
                        amplicon = assayInfo.Amplicon(sequence=_clean_seq(_strip(row[8].value)), variant_name=_clean_str(_strip(row[7].value)))
                        if element:
                            amplicon.add_SNP(element) if isinstance(element, assayInfo.SNP) else amplicon.add_ROI(element)
                        else:
                            amplicon.significance = significance
                elif amplicon and element:  # Continuing rows: attach another SNP/ROI to the current Amplicon
                    amplicon.add_SNP(element) if isinstance(element, assayInfo.SNP) else amplicon.add_ROI(element)

                if target and _strip(row[8].value):
                    target.add_amplicon(amplicon)
                    amplicon = None
                elif target:
                    target.amplicon = amplicon
                else:
                    target = assayInfo.Target(function=_strip(row[2].value), gene_name=_strip(row[3].value), start_position=_strip(row[4].value), end_position=_strip(row[5].value), reverse_comp=_strip(row[6].value), amplicon=amplicon)
                    assay.target = target

            if assay:
                assay_list.append(assay)

        # Write final JSON output
        assay_data = {"Assay": assay_list}
        assayInfo.writeJSON(assay_data, out_file)
        return 0

    except Exception as e:
        if DEBUG: raise e
        program_name = os.path.basename(sys.argv[0])
        sys.stderr.write(f"{program_name}: {repr(e)}\n")
        sys.stderr.write("  for help use --help\n")
        return 2

if __name__ == "__main__":
    sys.exit(main())